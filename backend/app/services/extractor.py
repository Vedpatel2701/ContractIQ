import os
from pathlib import Path
from pypdf import PdfReader
from docx import Document
from app.services.ocr import extract_text_from_image

def extract_text_from_file(file_path: str) -> str:
    """
    Extracts text from PDF, DOCX, TXT, or Image files.
    """
    path = Path(file_path)
    suffix = path.suffix.lower()
    
    if suffix == ".pdf":
        return extract_from_pdf(file_path)
    elif suffix in [".docx", ".doc"]:
        return extract_from_docx(file_path)
    elif suffix in [".txt", ".rtf", ".md"]:
        return extract_from_text(file_path)
    elif suffix in [".png", ".jpg", ".jpeg", ".tiff", ".bmp"]:
        ocr_text = extract_text_from_image(file_path)
        if ocr_text.strip():
            return f"--- Page 1 ---\n{ocr_text.strip()}"
        return ""
    else:
        # Fallback raw read
        return extract_from_text(file_path)

def extract_from_pdf(pdf_path: str) -> str:
    """
    Extracts text from PDF page by page.
    Uses native text extraction when sufficient text is found.
    When a page contains scanned content or insufficient extractable text (<40 chars),
    runs OCR specifically on that page to preserve page attribution.
    Reads PDF into BytesIO buffer so disk file handles are released immediately on Windows.
    """
    extracted_pages = []
    try:
        import io
        with open(pdf_path, "rb") as f:
            pdf_bytes = f.read()

        reader = PdfReader(io.BytesIO(pdf_bytes))
        total_pages = len(reader.pages)

        pdfium_doc = None
        try:
            import pypdfium2 as pdfium
            pdfium_doc = pdfium.PdfDocument(pdf_bytes)
        except Exception as pe_err:
            print(f"[PDF Extraction] pdfium renderer notice: {pe_err}")

        for i in range(total_pages):
            page_num = i + 1
            page_text = ""

            try:
                raw_text = reader.pages[i].extract_text() or ""
                clean_raw = raw_text.strip()
                # Use native text extraction if page has sufficient content (>= 40 chars)
                if len(clean_raw) >= 40:
                    page_text = clean_raw
            except Exception as pe:
                print(f"[PDF Extraction] Error reading text on page {page_num}: {pe}")

            # Per-page OCR fallback when native text extraction is missing or insufficient (< 40 chars)
            if not page_text and pdfium_doc:
                try:
                    rendered_page = pdfium_doc[i].render(scale=2).to_pil()
                    ocr_text = extract_text_from_image(rendered_page)
                    if ocr_text.strip():
                        page_text = ocr_text.strip()
                        print(f"[PDF Extraction] OCR successfully extracted text for scanned Page {page_num}")
                except Exception as oe:
                    print(f"[PDF Extraction] Page {page_num} OCR error: {oe}")

            if page_text:
                extracted_pages.append(f"--- Page {page_num} ---\n{page_text}")

        if pdfium_doc:
            try:
                pdfium_doc.close()
            except Exception:
                pass

        return "\n\n".join(extracted_pages)
    except Exception as e:
        print(f"[PDF Extraction Error] {e}")
        return ""

def extract_from_docx(docx_path: str) -> str:
    try:
        doc = Document(docx_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n\n".join(paragraphs)
    except Exception as e:
        print(f"[DOCX Extraction Error] {e}")
        return ""

def extract_from_text(text_path: str) -> str:
    try:
        with open(text_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    except Exception as e:
        print(f"[Text Read Error] {e}")
        return ""
